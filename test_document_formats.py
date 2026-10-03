from pathlib import Path
from tempfile import TemporaryDirectory

import pandas as pd
from docx import Document as DocxDocument
from openpyxl import Workbook
from pptx import Presentation

from src.document_loader import load_documents


# ================================================================
# TEST FIXTURE CREATION
# ================================================================

def create_docx(path: Path):
    """Create a sample DOCX file."""

    document = DocxDocument()

    document.add_heading(
        "Machine Learning",
        level=1,
    )

    document.add_paragraph(
        "Machine learning uses data to learn patterns "
        "and make predictions."
    )

    table = document.add_table(
        rows=2,
        cols=2,
    )

    table.cell(0, 0).text = "Model"
    table.cell(0, 1).text = "Accuracy"

    table.cell(1, 0).text = "Random Forest"
    table.cell(1, 1).text = "0.91"

    document.save(path)


def create_pptx(path: Path):
    """Create a sample PPTX file."""

    presentation = Presentation()

    slide = presentation.slides.add_slide(
        presentation.slide_layouts[1]
    )

    slide.shapes.title.text = "AI Engineering"

    slide.placeholders[1].text = (
        "Python\n"
        "Machine Learning\n"
        "RAG"
    )

    presentation.save(path)


def create_xlsx(path: Path):
    """Create a sample XLSX workbook with multiple sheets."""

    employees = pd.DataFrame(
        {
            "name": [
                "Alice",
                "Bob",
            ],
            "role": [
                "Data Scientist",
                "ML Engineer",
            ],
            "experience": [
                2,
                3,
            ],
        }
    )

    projects = pd.DataFrame(
        {
            "project": [
                "Churn Prediction",
                "RAG Assistant",
            ],
            "technology": [
                "Scikit-learn",
                "LangChain",
            ],
            "status": [
                "Completed",
                "In Progress",
            ],
        }
    )

    with pd.ExcelWriter(
        path,
        engine="openpyxl",
    ) as writer:

        employees.to_excel(
            writer,
            index=False,
            sheet_name="Employees",
        )

        projects.to_excel(
            writer,
            index=False,
            sheet_name="Projects",
        )


def create_xlsx_with_blank_rows(path: Path):
    """Create a sheet whose true row numbers shift around blanks."""

    frame = pd.DataFrame(
        {
            "name": [
                "Alice",
                None,
                "Bob",
                "Carl",
            ],
            "score": [
                2,
                None,
                3,
                None,
            ],
        }
    )

    with pd.ExcelWriter(
        path,
        engine="openpyxl",
    ) as writer:

        frame.to_excel(
            writer,
            index=False,
            sheet_name="Scores",
        )


def create_xlsx_with_blank_header(path: Path):
    """Create a sheet with an unnamed column and mixed value types."""

    frame = pd.DataFrame(
        {
            "name": [
                "Alice",
                "Bob",
            ],
            None: [
                "mid",
                "mid",
            ],
            "score": [
                2,
                3.5,
            ],
        }
    )

    with pd.ExcelWriter(
        path,
        engine="openpyxl",
    ) as writer:

        frame.to_excel(
            writer,
            index=False,
            sheet_name="Mixed",
        )


def create_empty_xlsx(path: Path):
    """Create a workbook with a single empty worksheet."""

    workbook = Workbook()

    workbook.active.title = "Empty"

    workbook.save(path)


# ================================================================
# FORMAT TESTS
# ================================================================

def test_txt(temp_dir: Path):
    """Test TXT ingestion."""

    path = temp_dir / "notes.txt"

    path.write_text(
        "Python is widely used for AI engineering.",
        encoding="utf-8",
    )

    documents = load_documents(path)

    assert len(documents) == 1

    assert "Python" in documents[0].page_content

    assert (
        documents[0].metadata["file_type"]
        == "txt"
    )

    assert (
        documents[0].metadata["document_name"]
        == "notes.txt"
    )


def test_markdown(temp_dir: Path):
    """Test Markdown ingestion."""

    path = temp_dir / "notes.md"

    path.write_text(
        "# Machine Learning\n\n"
        "Machine learning learns patterns from data.",
        encoding="utf-8",
    )

    documents = load_documents(path)

    assert len(documents) == 1

    assert (
        "Machine Learning"
        in documents[0].page_content
    )

    assert (
        documents[0].metadata["file_type"]
        == "markdown"
    )

    assert (
        documents[0].metadata["document_name"]
        == "notes.md"
    )


def test_csv(temp_dir: Path):
    """Test CSV ingestion."""

    path = temp_dir / "employees.csv"

    path.write_text(
        "name,role,experience\n"
        "Alice,Data Scientist,2\n"
        "Bob,ML Engineer,3\n",
        encoding="utf-8",
    )

    documents = load_documents(path)

    assert len(documents) == 2

    # First row
    assert (
        documents[0].metadata["file_type"]
        == "csv"
    )

    assert (
        documents[0].metadata["row"]
        == 2
    )

    assert (
        documents[0].metadata["document_name"]
        == "employees.csv"
    )

    assert (
        "name: Alice"
        in documents[0].page_content
    )

    assert (
        "role: Data Scientist"
        in documents[0].page_content
    )

    assert (
        "experience: 2"
        in documents[0].page_content
    )

    # Second row
    assert (
        documents[1].metadata["row"]
        == 3
    )

    assert (
        "name: Bob"
        in documents[1].page_content
    )


def test_xlsx(temp_dir: Path):
    """Test XLSX ingestion with multiple worksheets."""

    path = temp_dir / "employees.xlsx"

    create_xlsx(path)

    documents = load_documents(path)

    # Two rows from Employees +
    # two rows from Projects.
    assert len(documents) == 4

    # ------------------------------------------------------------
    # Employees sheet
    # ------------------------------------------------------------

    employee_documents = [
        document
        for document in documents
        if document.metadata["sheet"]
        == "Employees"
    ]

    assert len(employee_documents) == 2

    first_employee = employee_documents[0]

    assert (
        first_employee.metadata["file_type"]
        == "xlsx"
    )

    assert (
        first_employee.metadata["sheet"]
        == "Employees"
    )

    assert (
        first_employee.metadata["row"]
        == 2
    )

    assert (
        first_employee.metadata["document_name"]
        == "employees.xlsx"
    )

    assert (
        "name: Alice"
        in first_employee.page_content
    )

    assert (
        "role: Data Scientist"
        in first_employee.page_content
    )

    assert (
        "experience: 2"
        in first_employee.page_content
    )

    second_employee = employee_documents[1]

    assert (
        second_employee.metadata["row"]
        == 3
    )

    assert (
        "name: Bob"
        in second_employee.page_content
    )

    # ------------------------------------------------------------
    # Projects sheet
    # ------------------------------------------------------------

    project_documents = [
        document
        for document in documents
        if document.metadata["sheet"]
        == "Projects"
    ]

    assert len(project_documents) == 2

    first_project = project_documents[0]

    assert (
        first_project.metadata["file_type"]
        == "xlsx"
    )

    assert (
        first_project.metadata["sheet"]
        == "Projects"
    )

    assert (
        first_project.metadata["row"]
        == 2
    )

    assert (
        "project: Churn Prediction"
        in first_project.page_content
    )

    assert (
        "technology: Scikit-learn"
        in first_project.page_content
    )

    assert (
        "status: Completed"
        in first_project.page_content
    )


def test_xlsx_preserves_real_row_numbers(
    temp_dir: Path,
):
    """
    Worksheet row numbers must survive blank rows.

    Blank rows are skipped, but the reported row number must stay
    aligned with the real row in the worksheet.
    """

    path = temp_dir / "scores.xlsx"

    create_xlsx_with_blank_rows(path)

    documents = load_documents(path)

    # Three populated rows; the blank row is skipped.
    assert len(documents) == 3

    # Header is worksheet row 1, so Alice sits on row 2.
    assert documents[0].metadata["row"] == 2
    assert documents[0].metadata["sheet"] == "Scores"
    assert "name: Alice" in documents[0].page_content

    # Worksheet row 3 is blank, so Bob must be reported as row 4.
    assert documents[1].metadata["row"] == 4
    assert "name: Bob" in documents[1].page_content

    # Worksheet row 5 holds Carl.
    assert documents[2].metadata["row"] == 5
    assert "name: Carl" in documents[2].page_content

    # Row labels are written into the content as well.
    assert "Excel row: 4" in documents[1].page_content
    assert "Excel row: 5" in documents[2].page_content


def test_xlsx_normalizes_blank_column_names(
    temp_dir: Path,
):
    """Blank headers must not leak reader placeholders."""

    path = temp_dir / "mixed.xlsx"

    create_xlsx_with_blank_header(path)

    documents = load_documents(path)

    assert len(documents) == 2

    content = documents[0].page_content

    # The unnamed column becomes a positional name.
    assert "column_2: mid" in content

    # Internal reader naming never reaches the text.
    assert "Unnamed" not in content

    assert "name: Alice" in content

    # Integer cells avoid float noise.
    assert "score: 2" in content
    assert "score: 2.0" not in content

    assert "score: 3.5" in documents[1].page_content


def test_xlsx_empty_workbook_raises(
    temp_dir: Path,
):
    """A workbook with no data rows is rejected."""

    path = temp_dir / "empty.xlsx"

    create_empty_xlsx(path)

    try:
        load_documents(path)
    except ValueError as error:
        assert "No readable data" in str(error)
    else:
        raise AssertionError(
            "Empty workbook must raise ValueError."
        )


def test_xlsx_corrupt_file_raises(
    temp_dir: Path,
):
    """A file that is not a workbook is rejected."""

    path = temp_dir / "corrupt.xlsx"

    path.write_bytes(b"this is not an excel workbook")

    try:
        load_documents(path)
    except ValueError as error:
        assert (
            "Unable to read Excel workbook"
            in str(error)
        )
    else:
        raise AssertionError(
            "Corrupt workbook must raise ValueError."
        )


def test_docx(temp_dir: Path):
    """Test DOCX ingestion."""

    path = temp_dir / "document.docx"

    create_docx(path)

    documents = load_documents(path)

    assert len(documents) == 1

    assert (
        documents[0].metadata["file_type"]
        == "docx"
    )

    assert (
        documents[0].metadata["document_name"]
        == "document.docx"
    )

    assert (
        "Machine Learning"
        in documents[0].page_content
    )

    assert (
        "Random Forest"
        in documents[0].page_content
    )

    assert (
        "0.91"
        in documents[0].page_content
    )


def test_pptx(temp_dir: Path):
    """Test PPTX ingestion."""

    path = temp_dir / "presentation.pptx"

    create_pptx(path)

    documents = load_documents(path)

    assert len(documents) == 1

    assert (
        documents[0].metadata["file_type"]
        == "pptx"
    )

    assert (
        documents[0].metadata["slide"]
        == 1
    )

    assert (
        documents[0].metadata["document_name"]
        == "presentation.pptx"
    )

    assert (
        "AI Engineering"
        in documents[0].page_content
    )

    assert (
        "Python"
        in documents[0].page_content
    )

    assert (
        "Machine Learning"
        in documents[0].page_content
    )

    assert (
        "RAG"
        in documents[0].page_content
    )


# ================================================================
# TEST RUNNER
# ================================================================

def main():
    """Run all document format tests."""

    with TemporaryDirectory() as temp:

        temp_dir = Path(temp)

        # --------------------------------------------------------
        # TXT
        # --------------------------------------------------------

        test_txt(temp_dir)
        print("TXT: PASS")

        # --------------------------------------------------------
        # Markdown
        # --------------------------------------------------------

        test_markdown(temp_dir)
        print("Markdown: PASS")

        # --------------------------------------------------------
        # CSV
        # --------------------------------------------------------

        test_csv(temp_dir)
        print("CSV: PASS")

        # --------------------------------------------------------
        # XLSX
        # --------------------------------------------------------

        test_xlsx(temp_dir)
        print("XLSX: PASS")

        test_xlsx_preserves_real_row_numbers(
            temp_dir
        )
        print("XLSX row numbers: PASS")

        test_xlsx_normalizes_blank_column_names(
            temp_dir
        )
        print("XLSX column names: PASS")

        test_xlsx_empty_workbook_raises(
            temp_dir
        )
        print("XLSX empty workbook: PASS")

        test_xlsx_corrupt_file_raises(
            temp_dir
        )
        print("XLSX corrupt file: PASS")

        # --------------------------------------------------------
        # DOCX
        # --------------------------------------------------------

        test_docx(temp_dir)
        print("DOCX: PASS")

        # --------------------------------------------------------
        # PPTX
        # --------------------------------------------------------

        test_pptx(temp_dir)
        print("PPTX: PASS")

    print(
        "\nALL DOCUMENT FORMAT TESTS PASSED"
    )


if __name__ == "__main__":
    main()